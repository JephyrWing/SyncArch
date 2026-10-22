"""Step 2 기능 명세를 팀 공용 양식의 API 명세 엑셀로 만듭니다."""
import io
import os
import re
import threading
import uuid
from contextlib import suppress
from datetime import timedelta
from pathlib import Path

from flask import current_app
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..errors import ProjectNotFound
from ..repositories import project_repository as projects
from ..repositories.user_repository import find_by_id
from .project_access import require_access

# 팀 공용 양식입니다. 열 순서와 이름을 바꾸면 이 파일을 쓰는 팀원 문서가 깨지므로 함께 알립니다.
COLUMNS = [
    ("No", 6, None),
    ("기능 ID", 10, "id"),
    ("도메인", 12, "domain"),
    ("기능명", 24, "name"),
    ("우선순위", 9, "priority"),
    ("Method", 9, "method"),
    ("Endpoint", 32, "endpoint"),
    ("Request 필드", 36, "requestFields"),
    ("Response 필드", 36, "responseFields"),
    ("응답 코드", 14, "responseStatusCodes"),
    ("담당자", 10, "assignee"),
    ("진행 상태", 10, "status"),
    ("설명", 40, "summary"),
    ("연관 아이디어", 24, "relatedStoryId"),
]
STATUS_LABELS = {"PLANNED": "계획", "IN_PROGRESS": "진행중", "QA": "QA 검증", "COMPLETED": "완료됨"}
HEADER_ROW = 4
NAVY = "1F3A5F"
THIN = Side(style="thin", color="C5D3E8")
# 사람이 읽는 문서라 한국 시간으로 씁니다. 한국은 일광절약시간이 없어 고정 +9시간이면 됩니다.
KST = timedelta(hours=9)
# shortcut: 프로세스 안에서만 순서를 보장합니다. 서버 프로세스를 여러 개로 늘리면 파일 잠금이나 작업 큐로 바꿉니다.
_export_lock = threading.Lock()
_exported_versions = {}


def _text(cell, value):
    # "="로 시작하는 입력이 수식으로 실행되지 않게 항상 문자열로 저장합니다(수식 주입 방지).
    # 엑셀이 받지 않는 제어문자는 지웁니다. 남기면 파일 생성 자체가 실패합니다.
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def build_workbook(project, document, updater):
    """저장된 문서의 Step 2 기능 명세로 정해진 양식의 .xlsx 바이트를 만듭니다."""
    payload = document.payload
    ideas = {idea.get("id"): idea.get("title", "") for idea in payload.get("ideas", [])}
    book = Workbook()
    sheet = book.active
    sheet.title = "API 명세"

    _text(sheet.cell(1, 1), f"{project.title} API 명세서")
    sheet.cell(1, 1).font = Font(size=16, bold=True, color=NAVY)
    _text(sheet.cell(2, 1), f"문서 버전 v{document.version} · 마지막 저장 {(document.updated_at + KST):%Y-%m-%d %H:%M:%S} (KST) · "
                            f"{updater.display_name if updater else '-'} · SyncArch에서 자동 생성, 직접 고치지 마세요")
    sheet.cell(2, 1).font = Font(size=9, color="5B6B80")

    for index, (title, width, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(HEADER_ROW, index)
        _text(cell, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        sheet.column_dimensions[cell.column_letter].width = width

    for number, feature in enumerate(payload.get("features", []), start=1):
        row = HEADER_ROW + number
        for index, (_, _, key) in enumerate(COLUMNS, start=1):
            if key is None:
                value = number
            elif key == "id":
                # 계약 v2의 기능은 UUID id와 화면용 displayCode(FEAT-01)를 따로 가질 수 있습니다.
                value = feature.get("displayCode") or feature.get("id", "")
            elif key == "status":
                value = STATUS_LABELS.get(feature.get(key), feature.get(key, ""))
            elif key == "responseStatusCodes":
                value = ", ".join(feature.get(key) or [])
            elif key == "relatedStoryId":
                value = ideas.get(feature.get(key), "")
            else:
                value = feature.get(key, "")
            cell = sheet.cell(row, index)
            _text(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

    sheet.freeze_panes = sheet.cell(HEADER_ROW + 1, 1)
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{sheet.cell(HEADER_ROW, len(COLUMNS)).column_letter}{max(HEADER_ROW, sheet.max_row)}"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def file_name(project):
    # Windows 파일 이름에 못 쓰는 문자를 바꾸고, 같은 이름의 프로젝트끼리 섞이지 않게 ID 앞자리를 붙입니다.
    safe = re.sub(r'[\\/:*?"<>|\s]+', "_", project.title).strip("_.") or "project"
    return f"{safe}_{project.id[:8]}_API명세.xlsx"


def _workbook(session, project):
    document = projects.find_document(session, project.id)
    if document is None:
        raise ProjectNotFound()
    return document, build_workbook(project, document, find_by_id(session, document.updated_by))


def api_spec_file(session_factory, actor_id, project_id):
    """멤버만 받습니다. 항상 마지막으로 저장된 버전으로 만듭니다."""
    with session_factory() as session:
        project, _ = require_access(session, actor_id, project_id)
        return _workbook(session, project)[1], file_name(project)


def export_api_spec(session_factory, project_id):
    """API_SPEC_EXPORT_DIR이 있으면 저장할 때마다 그 폴더의 파일을 최신본으로 바꿉니다.

    폴더를 구글 드라이브 동기화 폴더로 두면 팀원 모두가 같은 최신 엑셀을 봅니다. 문서 저장은 이미 커밋됐으므로
    어떤 실패도 저장 요청을 실패시키지 않고 로그만 남깁니다(엑셀에서 파일을 열어 두면 Windows가 파일을 잠급니다).
    """
    folder = current_app.config.get("API_SPEC_EXPORT_DIR")
    if not folder:
        return
    try:
        # 저장이 겹쳐도 늦게 끝난 옛 버전이 최신 파일을 덮어쓰지 않게 순서대로 쓰고 버전을 비교합니다.
        with _export_lock, session_factory() as session:
            project = projects.find_project(session, project_id)
            document, content = _workbook(session, project)
            key = (folder, project_id)
            if _exported_versions.get(key, 0) >= document.version:
                return
            target = Path(folder) / file_name(project)
            temp = target.with_name(f"{target.name}.{uuid.uuid4().hex}.tmp")
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                # 제목이 바뀌면 옛 파일의 이름을 바꿔 이어 씁니다. 드라이브 공유 링크가 같은 파일을 계속 가리킵니다.
                for old in target.parent.glob(f"*_{project.id[:8]}_API명세.xlsx"):
                    if old != target and not target.exists():
                        os.replace(old, target)
                temp.write_bytes(content)
                os.replace(temp, target)
                _exported_versions[key] = document.version
            finally:
                with suppress(OSError):
                    temp.unlink(missing_ok=True)
    # 커밋 뒤의 부가 작업이라 어떤 예외든 저장 응답을 깨뜨리면 안 됩니다.
    except Exception as error:
        current_app.logger.warning("API spec export skipped (%s) for project %s", type(error).__name__, project_id)
