"""초대 코드·채팅·접속 표시·API 명세 엑셀."""
from contextlib import nullcontext
import io
from types import SimpleNamespace
import uuid
from unittest.mock import patch

from openpyxl import load_workbook

from test_auth import auth_app  # noqa: F401
from test_projects import add_editor, create, users  # noqa: F401


def save(users, project_id, base_version, payload, title="설계 프로젝트"):
    return users["owner"].put(f"/api/projects/{project_id}/document", headers=users["owner_headers"], json={
        "baseVersion": base_version, "mutationId": str(uuid.uuid4()),
        "project": {"title": title, "summary": "", "categoryTag": "", "targetDbms": "mariadb"}, "payload": payload})


IDEA_ID = "6f1d3c1e-4b7a-4c55-9a51-3f2a6c0b9d10"


def feature(code, name, **extra):
    # 계약 v2: id는 UUID이고 화면용 코드는 displayCode입니다. 같은 코드는 저장마다 같은 id를 씁니다.
    return {"id": str(uuid.uuid5(uuid.NAMESPACE_URL, code)), "displayCode": code, "domain": "인증", "name": name, "priority": "P0", "assignee": "앨리스", "method": "POST",
            "endpoint": "/api/auth/login", "requestFields": "loginId: string", "responseFields": "user: object",
            "summary": "로그인", "status": "IN_PROGRESS", **extra}


def spec_payload(*features, idea=False):
    # 서버가 아이디어에 authorId를 채우므로, 같은 아이디어를 다시 보내는 저장에는 그 값을 함께 보내야 합니다.
    ideas = [{"id": IDEA_ID, "category": "user_stories", "title": "구글로 로그인"}] if idea else []
    return {"schemaVersion": 2, "ideas": ideas,
            "features": list(features), "nodes": [], "connections": [], "tables": []}


def rows(content):
    sheet = load_workbook(io.BytesIO(content)).active
    return sheet, [[cell.value for cell in row] for row in sheet.iter_rows(min_row=4)]


def test_owner_code_lets_a_teammate_join_as_editor(users):
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    code = owner.get(f"/api/projects/{project_id}/invite", headers=oh).json["data"]["code"]
    assert len(code) == 8 and code.isalnum() and code.isupper()
    # 멤버가 아니면 코드를 볼 수 없습니다.
    assert editor.get(f"/api/projects/{project_id}/invite", headers=eh).status_code == 404

    joined = editor.post("/api/projects/join", json={"code": f"{code[:4].lower()}-{code[4:]} "}, headers=eh)
    assert joined.status_code == 200
    assert joined.json["data"]["id"] == project_id and joined.json["data"]["role"] == "EDITOR"
    # 참가는 멤버 추가와 같이 문서 버전을 올리고 이력을 남깁니다.
    assert joined.json["data"]["version"] == 2
    assert editor.post("/api/projects/join", json={"code": code}, headers=eh).json["data"]["version"] == 2
    # 편집자는 코드를 볼 수 없습니다.
    assert editor.get(f"/api/projects/{project_id}/invite", headers=eh).status_code == 403
    roles = {m["displayName"]: m["role"] for m in owner.get(f"/api/projects/{project_id}/members", headers=oh).json["data"]}
    assert roles == {"검증 A": "OWNER", "검증 B": "EDITOR"}
    activity = owner.get(f"/api/projects/{project_id}/activity", headers=oh).json["data"]
    assert "member_added" in str(activity)


def test_regenerated_code_replaces_the_old_one_and_bad_codes_fail(users):
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    old = owner.get(f"/api/projects/{project_id}/invite", headers=oh).json["data"]["code"]
    new = owner.post(f"/api/projects/{project_id}/invite", headers=oh).json["data"]["code"]
    assert new != old
    assert editor.post("/api/projects/join", json={"code": old}, headers=eh).status_code == 404
    assert editor.post("/api/projects/join", json={"code": " - "}, headers=eh).status_code == 400
    assert editor.post("/api/projects/join", json={"code": new}).status_code == 401
    owner.patch(f"/api/projects/{project_id}/archive", json={"archived": True}, headers=oh)
    assert editor.post("/api/projects/join", json={"code": new}, headers=eh).json["error"]["code"] == "PROJECT_ARCHIVED"


def test_chat_is_shared_by_members_and_supports_incremental_polling(users):
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    assert editor.post(f"/api/projects/{project_id}/messages", json={"body": "몰래"}, headers=eh).status_code == 404
    code = owner.get(f"/api/projects/{project_id}/invite", headers=oh).json["data"]["code"]
    editor.post("/api/projects/join", json={"code": code}, headers=eh)

    first = owner.post(f"/api/projects/{project_id}/messages", json={"body": "  API 서버랑 DB 연결해 주세요 "}, headers=oh)
    assert first.status_code == 201 and first.json["data"]["body"] == "API 서버랑 DB 연결해 주세요"
    editor.post(f"/api/projects/{project_id}/messages", json={"body": "넵 SQL로 연결할게요"}, headers=eh)
    history = editor.get(f"/api/projects/{project_id}/messages", headers=eh).json["data"]
    assert [(m["displayName"], m["body"]) for m in history] == [("검증 A", "API 서버랑 DB 연결해 주세요"), ("검증 B", "넵 SQL로 연결할게요")]

    since = history[-1]["createdAt"]
    owner.post(f"/api/projects/{project_id}/messages", json={"body": "확인!"}, headers=oh)
    newer = [m["body"] for m in owner.get(f"/api/projects/{project_id}/messages", query_string={"since": since}, headers=oh).json["data"]]
    # 같은 시각 메시지를 놓치지 않으려고 since 시각 메시지도 포함합니다(클라이언트가 id로 거릅니다).
    assert newer[-2:] == ["넵 SQL로 연결할게요", "확인!"] and len(newer) <= 3

    assert owner.post(f"/api/projects/{project_id}/messages", json={"body": "   "}, headers=oh).status_code == 400
    assert owner.post(f"/api/projects/{project_id}/messages", json={"body": "가" * 1001}, headers=oh).status_code == 400
    assert owner.get(f"/api/projects/{project_id}/messages", query_string={"since": "어제"}, headers=oh).status_code == 400
    owner.patch(f"/api/projects/{project_id}/archive", json={"archived": True}, headers=oh)
    assert editor.post(f"/api/projects/{project_id}/messages", json={"body": "보관 후"}, headers=eh).status_code == 403


def test_chat_cursor_reads_past_more_than_a_page_of_same_time_messages(users):
    from datetime import datetime
    from app.extensions import db
    from app.models import ProjectMessage
    owner, oh = users["owner"], users["owner_headers"]
    project_id = create(users)["id"]
    same = datetime(2026, 10, 22, 1, 0, 0, 123456)
    with users["app"].app_context():
        db.session.add_all(ProjectMessage(project_id=project_id, user_id=users["owner_id"], body=f"m{i}", created_at=same)
                           for i in range(205))
        db.session.add(ProjectMessage(project_id=project_id, user_id=users["owner_id"], body="later",
                                      created_at=datetime(2026, 10, 22, 1, 0, 1)))
        db.session.commit()

    def page(since, after=None):
        query = {"since": since} | ({"after": after} if after else {})
        response = owner.get(f"/api/projects/{project_id}/messages", query_string=query, headers=oh)
        assert response.status_code == 200
        return response.json["data"]

    first = page("2026-10-22T01:00:00Z")
    assert len(first) == 200 and len({m["createdAt"] for m in first}) == 1
    assert [m["id"] for m in first] == sorted(m["id"] for m in first)
    # 예전 방식(since만)은 같은 시각 200개를 다시 받아 멈추지만, (시각, id) 커서는 다음 메시지로 넘어갑니다.
    assert page(first[-1]["createdAt"]) == first
    second = page(first[-1]["createdAt"], first[-1]["id"])
    assert [m["body"] for m in second][-1] == "later" and len(second) == 6
    assert {m["id"] for m in first}.isdisjoint(m["id"] for m in second)
    assert page(second[-1]["createdAt"], second[-1]["id"]) == []

    assert owner.get(f"/api/projects/{project_id}/messages", query_string={"after": first[0]["id"]},
                     headers=oh).status_code == 400
    assert owner.get(f"/api/projects/{project_id}/messages", query_string={"since": first[0]["createdAt"], "after": "x"},
                     headers=oh).status_code == 400


def test_resending_a_chat_message_with_the_same_client_id_saves_it_once(users):
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    add_editor(users, project_id)
    url = f"/api/projects/{project_id}/messages"
    key = str(uuid.uuid4())

    first = owner.post(url, json={"body": "한 번만", "clientId": key}, headers=oh)
    # 응답을 잃은 브라우저가 같은 키로 다시 보내면 새로 만들지 않고 처음 메시지를 돌려줍니다.
    again = owner.post(url, json={"body": " 한 번만 ", "clientId": key}, headers=oh)
    assert first.status_code == 201 and again.status_code == 200 and again.json["data"] == first.json["data"]
    # 키는 보낸 사람마다 따로라 다른 멤버가 같은 키를 써도 섞이지 않습니다.
    assert editor.post(url, json={"body": "한 번만", "clientId": key}, headers=eh).status_code == 201
    reused = owner.post(url, json={"body": "다른 내용", "clientId": key}, headers=oh)
    assert reused.status_code == 409 and reused.json["error"]["code"] == "IDEMPOTENCY_KEY_REUSED"
    assert owner.post(url, json={"body": "키 없음"}, headers=oh).status_code == 201
    assert owner.post(url, json={"body": "키 없음"}, headers=oh).status_code == 201
    assert owner.post(url, json={"body": "x", "clientId": "not-a-uuid"}, headers=oh).status_code == 400
    bodies = [m["body"] for m in owner.get(url, headers=oh).json["data"]]
    assert bodies == ["한 번만", "한 번만", "키 없음", "키 없음"]

    # 같은 키 두 요청이 동시에 확인을 통과한 경우: 늦은 쪽은 고유 인덱스에 막히고 먼저 저장된 메시지를 받습니다.
    from app.repositories import message_repository
    real = message_repository.find_by_client_id
    calls = []
    with patch.object(message_repository, "find_by_client_id",
                      lambda *args: calls.append(1) or (None if len(calls) == 1 else real(*args))):
        raced = owner.post(url, json={"body": "한 번만", "clientId": key}, headers=oh)
    assert raced.status_code == 200 and raced.json["data"] == first.json["data"] and len(calls) == 2
    assert len(owner.get(url, headers=oh).json["data"]) == 4


def test_presence_lists_members_with_their_step_and_expires(users, monkeypatch):
    clock = [1000.0]
    monkeypatch.setattr("app.services.presence_service.time.monotonic", lambda: clock[0])
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    owner.post(f"/api/projects/{project_id}/members", json={"loginId": "auth-b"}, headers=oh)

    assert owner.post(f"/api/projects/{project_id}/presence", json={"step": 3}, headers=oh).json["data"] == [
        {"userId": users["owner_id"], "displayName": "검증 A", "step": 3}]
    online = editor.post(f"/api/projects/{project_id}/presence", json={"step": 1}, headers=eh).json["data"]
    assert sorted((e["displayName"], e["step"]) for e in online) == [("검증 A", 3), ("검증 B", 1)]

    # 신호가 끊긴 멤버는 TTL 뒤 목록에서 빠지고, 나가기 요청은 바로 뺍니다.
    clock[0] += 16
    assert [e["displayName"] for e in editor.post(f"/api/projects/{project_id}/presence", json={"step": 2}, headers=eh).json["data"]] == ["검증 B"]
    assert editor.delete(f"/api/projects/{project_id}/presence", headers=eh).status_code == 204
    assert [e["displayName"] for e in owner.post(f"/api/projects/{project_id}/presence", json={"step": 3}, headers=oh).json["data"]] == ["검증 A"]


def test_presence_keeps_other_tabs_and_drops_removed_members(users):
    owner, editor = users["owner"], users["editor"]
    oh, eh = users["owner_headers"], users["editor_headers"]
    project_id = create(users)["id"]
    url = f"/api/projects/{project_id}/presence"
    owner.post(f"/api/projects/{project_id}/members", json={"loginId": "auth-b"}, headers=oh)

    # 같은 사람이 탭 두 개를 열면 한 줄로 보이고, 가장 최근 탭의 단계를 보여 줍니다.
    editor.post(url, json={"step": 1, "tabId": "tab-1"}, headers=eh)
    online = editor.post(url, json={"step": 4, "tabId": "tab-2"}, headers=eh).json["data"]
    assert [(e["displayName"], e["step"]) for e in online] == [("검증 B", 4)]

    # 한 탭을 닫아도 다른 탭은 접속 중으로 남습니다.
    assert editor.delete(f"{url}?tabId=tab-2", headers=eh).status_code == 204
    assert sorted((e["displayName"], e["step"]) for e in owner.post(url, json={"step": 2}, headers=oh).json["data"]) == [
        ("검증 A", 2), ("검증 B", 1)]

    # 멤버에서 빠진 사람은 TTL을 기다리지 않고 바로 목록에서 빠집니다.
    owner.delete(f'/api/projects/{project_id}/members/{users["editor_id"]}', headers=oh)
    assert [e["displayName"] for e in owner.post(url, json={"step": 2}, headers=oh).json["data"]] == ["검증 A"]
    assert owner.post(url, json={"step": 2, "tabId": ""}, headers=oh).status_code == 400


def test_presence_does_not_bring_back_a_member_removed_during_the_access_check(monkeypatch):
    from app.services import presence_service, project_events

    def removed_while_checking(session, actor_id, project_id):
        project_events.access_revoked.send("project-service", project_id=project_id, user_id=actor_id,
                                           reason="MEMBER_REMOVED")

    monkeypatch.setattr(presence_service, "require_access", removed_while_checking)
    monkeypatch.setattr(presence_service, "require_active_actor", lambda session, actor_id: SimpleNamespace(display_name="B"))
    assert presence_service.heartbeat(nullcontext, "user-b", "project-race", 1, "tab-1") == []

    # 권한 확인 뒤 TTL보다 오래 멈춘 신호는 제외 기록이 정리된 뒤라도 올리지 않습니다.
    clock = [1000.0]
    monkeypatch.setattr("app.services.presence_service.time.monotonic", lambda: clock[0])
    monkeypatch.setattr(presence_service, "require_access", lambda session, actor_id, project_id: clock.__setitem__(0, 1016.0))
    assert presence_service.heartbeat(nullcontext, "user-c", "project-slow", 1, "tab-1") == []


def test_presence_requires_membership_and_a_valid_step(users):
    project_id = create(users)["id"]
    assert users["outsider"].post(f"/api/projects/{project_id}/presence", json={"step": 1},
                                  headers=users["outsider_headers"]).status_code == 404
    for step in (0, 5, True, "1", None):
        response = users["owner"].post(f"/api/projects/{project_id}/presence", json={"step": step}, headers=users["owner_headers"])
        assert response.status_code == 400, step
    assert users["owner"].post(f"/api/projects/{project_id}/presence", json={"step": 1}).status_code == 401


def test_api_spec_download_uses_the_team_layout_and_latest_version(users):
    project_id = create(users)["id"]
    injected = feature("FEAT-02", "=HYPERLINK(\"http://evil\")", responseStatusCodes=["200", "401"],
                       relatedStoryId=IDEA_ID, status="COMPLETED", summary="줄\x01바꿈")
    response = save(users, project_id, 1, spec_payload(feature("FEAT-01", "로그인"), injected, idea=True))
    assert response.status_code == 200, response.json

    response = users["owner"].get(f"/api/projects/{project_id}/api-spec.xlsx", headers=users["owner_headers"])
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    assert "API" in response.headers["Content-Disposition"]
    sheet, table = rows(response.data)
    assert sheet.title == "API 명세" and sheet["A1"].value == "설계 프로젝트 API 명세서"
    assert sheet["A2"].value.startswith("문서 버전 v2")
    assert table[0][:4] == ["No", "기능 ID", "도메인", "기능명"]
    assert table[1][:4] == [1, "FEAT-01", "인증", "로그인"] and table[1][11] == "진행중"
    assert table[2][1] == "FEAT-02" and table[2][9:12] == ["200, 401", "앨리스", "완료됨"] and table[2][13] == "구글로 로그인"
    # 사용자가 입력한 "=..."는 수식이 아니라 글자로 남고, 엑셀이 받지 않는 제어문자는 빠집니다.
    assert table[2][3] == "=HYPERLINK(\"http://evil\")" and sheet["D6"].data_type == "s"
    assert table[2][12] == "줄바꿈"
    assert users["outsider"].get(f"/api/projects/{project_id}/api-spec.xlsx", headers=users["outsider_headers"]).status_code == 404


def test_saving_keeps_the_exported_file_current_even_when_export_fails(users, tmp_path):
    users["app"].config["API_SPEC_EXPORT_DIR"] = str(tmp_path)
    project_id = create(users)["id"]
    assert save(users, project_id, 1, spec_payload(feature("FEAT-01", "로그인"))).status_code == 200
    [exported] = list(tmp_path.glob("*.xlsx"))
    assert exported.name == f"설계_프로젝트_{project_id[:8]}_API명세.xlsx"

    assert save(users, project_id, 2, spec_payload(feature("FEAT-01", "로그인"), feature("FEAT-02", "로그아웃"))).status_code == 200
    sheet, table = rows(exported.read_bytes())
    assert sheet["A2"].value.startswith("문서 버전 v3") and table[2][3] == "로그아웃"

    # 누가 엑셀에서 파일을 열어 잠가도 문서 저장은 성공하고 파일만 다음 저장 때 갱신됩니다.
    with patch("app.services.api_spec_service.os.replace", side_effect=PermissionError):
        assert save(users, project_id, 3, spec_payload(feature("FEAT-01", "로그인"))).status_code == 200
    assert list(tmp_path.glob("*.tmp")) == []
    with patch("app.services.api_spec_service.build_workbook", side_effect=RuntimeError("boom")):
        assert save(users, project_id, 4, spec_payload(feature("FEAT-01", "로그인"))).status_code == 200
    assert save(users, project_id, 5, spec_payload(feature("FEAT-01", "로그인"))).status_code == 200
    assert rows(exported.read_bytes())[0]["A2"].value.startswith("문서 버전 v6")

    # 제목이 바뀌면 같은 파일의 이름을 바꿔 이어 씁니다.
    assert save(users, project_id, 6, spec_payload(feature("FEAT-01", "로그인")), title="새 이름").status_code == 200
    assert [path.name for path in tmp_path.glob("*.xlsx")] == [f"새_이름_{project_id[:8]}_API명세.xlsx"]
